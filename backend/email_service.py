"""
Core email logic shared by the Flask API.

Three jobs:
  1. make_fit_to_screen(html)  -> force any uploaded HTML into the same
     "fit-to-screen, no mobile reflow" shape we send by hand.
  2. load_recipients(...)       -> read col 3 (email) / col 4 (to|cc|bcc)
     from any similar .xlsx.
  3. send_email(...)            -> deliver via Office365 SMTP.
"""

import io
import re
import smtplib
from email.message import EmailMessage
from email.utils import make_msgid, formatdate

import openpyxl

# ---- SMTP / message defaults --------------------------------------------

SMTP_HOST = "smtp.office365.com"
SMTP_PORT = 587
DEFAULT_SENDER = "nudge-app@indegene.com"
DEFAULT_SUBJECT = "Account Intelligence Update"

# Which columns hold the data (1-based, matching the sheet).
EMAIL_COL = 3          # column 3 -> email address
DISPOSITION_COL = 4    # column 4 -> to / cc / bcc
VALID_DISPOSITIONS = {"to", "cc", "bcc"}


# =========================================================================
# 1. HTML transform: fit-to-screen, no mobile reflow
# =========================================================================

def _strip_maxwidth_media_blocks(html):
    """Remove every `@media ... { ... }` block whose query mentions
    max-width (the mobile reflow rules that distort the layout).

    Brace-balanced so nested braces inside the block are handled.
    Returns (new_html, removed_count).
    """
    out = []
    removed = 0
    i = 0
    lower = html.lower()
    while True:
        at = lower.find("@media", i)
        if at == -1:
            out.append(html[i:])
            break

        brace = html.find("{", at)
        if brace == -1:
            out.append(html[i:])
            break

        query = html[at:brace].lower()

        # Walk to the matching close brace.
        depth = 0
        j = brace
        while j < len(html):
            c = html[j]
            if c == "{":
                depth += 1
            elif c == "}":
                depth -= 1
                if depth == 0:
                    j += 1
                    break
            j += 1

        if "max-width" in query:
            # Drop the whole block (keep everything before it).
            out.append(html[i:at])
            removed += 1
            i = j
        else:
            # Keep this block untouched.
            out.append(html[i:j])
            i = j

    return "".join(out), removed


def apply_tagline(html, tagline):
    """Set the hero-band byline to exactly the text the user provides.

    - If a byline already exists ("For ... portfolio"), replace its text.
    - Otherwise insert one right after the "Portfolio intelligence · ..."
      banner line, styled to match (white, bold) on the dark band.

    The text is used verbatim — nothing is added or formatted.
    Returns (new_html, change_or_None).
    """
    byline_text = (tagline or "").strip()
    if not byline_text:
        return html, None

    # Escape markup-breaking characters; keep quotes/apostrophes as typed.
    safe = (
        byline_text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )

    # 1. Replace an existing byline if present. Lambda replacement avoids
    #    backreference interpretation of characters in the user's text.
    new_html, n = re.subn(
        r"For\s+[^<>]{1,120}?portfolio",
        lambda _m: safe, html, count=1, flags=re.I,
    )
    if n:
        return new_html, f"Set portfolio tagline to “{byline_text}”"

    # 2. No byline yet — insert one after the banner paragraph.
    byline_p = (
        '<p style="margin:0 0 18px 0;font-size:18px;line-height:24px;'
        'font-weight:800;color:#FFFFFF;">' + safe + "</p>"
    )

    banner = re.search(
        r"<p[^>]*>[^<]*?(?:portfolio intelligence|verified signals|signals)"
        r"[^<]*?</p>",
        html, flags=re.I,
    )
    if banner:
        at = banner.end()
        html = html[:at] + "\n    " + byline_p + html[at:]
        return html, f"Inserted portfolio tagline “{byline_text}”"

    # 3. Fallback — put it just before the first headline.
    h1 = re.search(r"<h1\b", html, flags=re.I)
    if h1:
        at = h1.start()
        html = html[:at] + byline_p + "\n    " + html[at:]
        return html, f"Inserted portfolio tagline “{byline_text}”"

    return html, None


def add_outlook_bgcolors(html):
    """Outlook (Word engine + OWA) unpredictably drops the CSS `background:`
    shorthand, so navy/colored bands render white. Add a matching legacy
    `bgcolor="#hex"` attribute — which Outlook always honors — to every
    table cell/table that sets a solid background color in its style.

    Browsers ignore bgcolor when the CSS works, so this is purely additive
    and safe. Returns (new_html, count_added).
    """
    added = 0
    tag_re = re.compile(r"<(td|th|tr|table)\b([^>]*)>", re.I)

    def repl(m):
        nonlocal added
        tag, attrs = m.group(1), m.group(2)
        if re.search(r"\bbgcolor\s*=", attrs, re.I):
            return m.group(0)  # already has one
        color_match = re.search(
            r"background(?:-color)?\s*:\s*(#[0-9a-fA-F]{3,6})", attrs, re.I
        )
        if not color_match:
            return m.group(0)
        added += 1
        return f'<{tag} bgcolor="{color_match.group(1)}"{attrs}>'

    return tag_re.sub(repl, html), added


def make_fit_to_screen(html):
    """Normalize any uploaded HTML to the exact shape we send:
      - viewport locked to a fixed 640px (client scales down to fit)
      - Apple's message-reformatting block removed so iOS Mail can scale
      - mobile reflow @media blocks stripped (they cause the distortion)

    Idempotent: running it on already-normalized HTML changes nothing.
    Returns (new_html, changes) where changes is a list of human-readable
    strings describing what was adjusted.
    """
    changes = []

    # 1. Remove x-apple-disable-message-reformatting meta (blocks auto-scale).
    html, n = re.subn(
        r'[ \t]*<meta[^>]*x-apple-disable-message-reformatting[^>]*>\s*\n?',
        "", html, flags=re.I,
    )
    if n:
        changes.append("Removed x-apple-disable-message-reformatting meta")

    # 2. Force viewport to width=640 (replace existing, or insert into head).
    viewport_tag = '<meta name="viewport" content="width=640">'
    if re.search(r'<meta[^>]+name=["\']viewport["\']', html, re.I):
        new_html, n = re.subn(
            r'<meta[^>]+name=["\']viewport["\'][^>]*>',
            viewport_tag, html, flags=re.I,
        )
        # Only count it as a change if the tag actually differed.
        if new_html != html:
            changes.append("Set viewport to fixed width=640")
        html = new_html
    else:
        if re.search(r"<head[^>]*>", html, re.I):
            html = re.sub(
                r"(<head[^>]*>)", r"\1\n" + viewport_tag, html,
                count=1, flags=re.I,
            )
        else:
            # No head at all — prepend one.
            html = "<head>\n" + viewport_tag + "\n</head>\n" + html
        changes.append("Inserted fixed viewport width=640")

    # 3. Strip mobile reflow @media blocks.
    html, removed = _strip_maxwidth_media_blocks(html)
    if removed:
        changes.append(
            f"Removed {removed} mobile reflow @media block"
            + ("s" if removed != 1 else "")
        )

    # 4. Add Outlook bgcolor fallbacks so colored bands don't render white.
    html, added = add_outlook_bgcolors(html)
    if added:
        changes.append(f"Added Outlook color fallbacks to {added} element(s)")

    if not changes:
        changes.append("Already fit-to-screen — no changes needed")

    return html, changes


def prepare_html(html, tagline=None):
    """Full pipeline applied to every uploaded HTML before preview/send:
    set the portfolio tagline (if given), then force fit-to-screen.
    Returns (new_html, changes).
    """
    changes = []
    html, tag_change = apply_tagline(html, tagline)
    if tag_change:
        changes.append(tag_change)
    html, fit_changes = make_fit_to_screen(html)
    changes.extend(fit_changes)
    return html, changes


# =========================================================================
# 2. Recipient loading
# =========================================================================

def _rows_to_buckets(rows):
    buckets = {"to": [], "cc": [], "bcc": []}
    for row in rows:
        if len(row) < max(EMAIL_COL, DISPOSITION_COL):
            continue
        email = row[EMAIL_COL - 1]
        disposition = row[DISPOSITION_COL - 1]
        if email is None or disposition is None:
            continue
        email = str(email).strip()
        disposition = str(disposition).strip().lower()
        if "@" not in email or disposition not in VALID_DISPOSITIONS:
            continue
        buckets[disposition].append(email)
    return buckets


def load_recipients_from_path(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    return _rows_to_buckets(ws.iter_rows(values_only=True))


def load_recipients_from_bytes(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    return _rows_to_buckets(ws.iter_rows(values_only=True))


# =========================================================================
# 3. Sending
# =========================================================================

def build_message(html, sender, buckets, subject):
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = sender
    if buckets["to"]:
        msg["To"] = ", ".join(buckets["to"])
    if buckets["cc"]:
        msg["Cc"] = ", ".join(buckets["cc"])
    # Bcc deliberately NOT a header — that's what keeps it blind.
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid(domain="indegene.com")
    msg.set_content("This message is best viewed in an HTML-capable email client.")
    msg.add_alternative(html, subtype="html")
    return msg


def send_email(html, buckets, subject, sender, smtp_user, smtp_pass):
    """Send the (already transformed) HTML to the recipient buckets.

    Returns the full envelope list actually delivered to.
    Raises smtplib exceptions on failure — the caller maps them to JSON.
    """
    all_recipients = buckets["to"] + buckets["cc"] + buckets["bcc"]
    if not all_recipients:
        raise ValueError("No valid recipients found.")
    if not buckets["to"]:
        raise ValueError("Sheet has no 'to' recipient. At least one is required.")

    msg = build_message(html, sender, buckets, subject)

    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=60) as server:
        server.ehlo()
        server.starttls()
        server.ehlo()
        server.login(smtp_user, smtp_pass)
        server.send_message(msg, to_addrs=all_recipients)

    return all_recipients
